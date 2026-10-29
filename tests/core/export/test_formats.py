"""各导出格式能被对应的程序读回来，内容不丢。"""

from __future__ import annotations

import io
import re
import zipfile
from xml.etree import ElementTree

from docx import Document as Docx

from roundtable.core.export import FORMATS, Style, filename, pdf, pptx
from roundtable.core.export.deliverable import Document

MD = """# 复习计划

> 怎么复习？

概要：**早**开始。

## 第一周

1. 列科目
2. 排时间

| 科目 | 时长 |
|---|---|
| 数学 | 3 小时 |

## 第二周 <script>

```python
    print("缩进")
```
"""
DOC = Document("复习计划", "怎么复习？", MD)


def test_docx_reads_back():
    d = Docx(io.BytesIO(FORMATS["docx"].render(DOC, Style())))
    texts = [p.text for p in d.paragraphs]
    assert "复习计划" in texts and "列科目" in texts and "第二周 <script>" in texts
    assert any(p.style.name.startswith("List Number") for p in d.paragraphs)
    assert d.tables[0].cell(1, 0).text == "数学"


def test_pptx_is_a_valid_package_with_one_slide_per_section():
    data = FORMATS["pptx"].render(DOC, Style())
    z = zipfile.ZipFile(io.BytesIO(data))
    for name in z.namelist():
        if name.endswith((".xml", ".rels")):
            ElementTree.fromstring(z.read(name))  # 每个部件都是合法 XML
    slides = sorted(n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
    assert len(slides) == 4  # 标题页、概要、第一周、第二周
    types = z.read("[Content_Types].xml").decode()
    assert all(f"/{s}" in types for s in slides)
    cover = z.read("ppt/slides/slide1.xml").decode()
    assert "复习计划" in cover and "怎么复习？" in cover
    assert "第二周 &lt;script&gt;" in z.read("ppt/slides/slide4.xml").decode()


def test_pptx_long_sections_continue_on_new_slides():
    md = "# T\n\n## 长\n\n" + "\n".join(f"- 第 {i} 条" for i in range(40))
    names = [name for name, _ in pptx.slides_of(md)]
    assert names[0] == "T" and names[1] == "长" and names[2] == "长（续）" and len(names) >= 4


def test_pdf_structure_and_pagination():
    data = FORMATS["pdf"].render(DOC, Style())
    assert data.startswith(b"%PDF-1.4") and data.rstrip().endswith(b"%%EOF")
    xref = int(data.rsplit(b"startxref\n", 1)[1].split(b"\n")[0])
    assert data[xref:].startswith(b"xref")
    table = data[xref:].split(b"\n")
    count = int(table[1].split()[1])
    for i in range(1, count):  # 每个对象的偏移都指对
        off = int(table[2 + i].split()[0])
        assert data[off:].startswith(f"{i} 0 obj".encode())
    assert b"<7F29> Tj <8FDB> Tj" in data  # 「缩进」逐字排版
    long = "\n\n".join(f"第 {i} 段，" + "很长的句子" * 30 for i in range(60))
    assert len(pdf.layout(long)) > 3
    from pypdf import PdfReader

    assert len(PdfReader(io.BytesIO(pdf.render(long, "长"))).pages) == len(pdf.layout(long))


def test_html_escapes_and_markdown_is_utf8():
    page = FORMATS["html"].render(DOC, Style()).decode()
    assert "<script>" not in page and "&lt;script&gt;" in page
    assert "<ol>" in page and "<table>" in page and "<strong>早</strong>" in page
    assert FORMATS["md"].render(DOC, Style()).decode() == MD


def test_filename_drops_unsafe_characters():
    doc = Document('a/b:c*"题目"<>|', "", "")
    assert filename(doc, "0123456789", "docx") == "圆桌-a_b_c__题目-01234567.docx"
