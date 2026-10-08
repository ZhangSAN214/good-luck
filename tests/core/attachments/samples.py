"""测试用的小文件（在内存中生成，不依赖仓库里的二进制文件）。"""

from __future__ import annotations

import io

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
JPG = b"\xff\xd8\xff\xe0" + b"\x00" * 64
MP3 = b"ID3\x04\x00" + b"\x00" * 64
WAV = b"RIFF\x24\x00\x00\x00WAVEfmt " + b"\x00" * 64


def pdf(*pages: str) -> bytes:
    """最小的带文字层的 PDF：每页一行文字（空字符串为没有文字的页，模拟扫描件）。"""
    objects: list[bytes] = []
    kids = []
    n_pages = len(pages)
    # 1 catalog, 2 pages, 3 font, 之后每页两个对象（page + content）
    for i in range(n_pages):
        kids.append(f"{4 + 2 * i} 0 R")
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {n_pages} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        content_no = 5 + 2 * i
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
                f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_no} 0 R >>"
            ).encode()
        )
        stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    import docx as python_docx

    document = python_docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
