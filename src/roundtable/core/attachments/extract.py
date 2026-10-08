"""从文档中提取文字：PDF（pypdf）、Word .docx（python-docx）、文本（自动识别编码）。

只读取文字，不执行文件中的任何内容。
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field

from roundtable.core.config.schema import UploadRules

from .detect import FileType, UploadError

# pypdf 遇到不规范的文件会大量输出警告，内容可能包含文件片段：不进入日志
logging.getLogger("pypdf").setLevel(logging.ERROR)


@dataclass(frozen=True)
class Extracted:
    text: str
    pages: int | None = None
    warnings: tuple[str, ...] = field(default_factory=tuple)


def _limit(text: str, rules: UploadRules, warnings: list[str]) -> str:
    text = text.strip()
    if len(text) > rules.max_text_chars:
        warnings.append(f"文字过长，只保留前 {rules.max_text_chars} 字")
        text = text[: rules.max_text_chars] + "\n……（以下内容过长，已截断）"
    return text


def extract_pdf(data: bytes, rules: UploadRules) -> Extracted:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UploadError("PDF 已加密，请先去掉密码")
        pages = len(reader.pages)
        if pages > rules.max_pdf_pages:
            raise UploadError(f"PDF 有 {pages} 页，超过上限 {rules.max_pdf_pages} 页")
        texts = [(page.extract_text() or "").strip() for page in reader.pages]
    except UploadError:
        raise
    except (PyPdfError, ValueError, KeyError, TypeError, OSError):
        raise UploadError("PDF 文件已损坏或无法读取") from None
    warnings: list[str] = []
    body = "\n\n".join(f"[第 {i} 页]\n{t}" for i, t in enumerate(texts, 1) if t)
    if sum(len(t) for t in texts) < rules.scanned_chars_per_page * max(pages, 1):
        warnings.append("这个 PDF 几乎没有文字层，可能是扫描件；目前不做文字识别，可以改传截图")
    return Extracted(_limit(body, rules, warnings), pages, tuple(warnings))


def extract_docx(data: bytes, rules: UploadRules) -> Extracted:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError, OSError):
        raise UploadError("Word 文件已损坏或无法读取") from None
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:  # 表格按行输出，单元格用 | 分隔
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    warnings: list[str] = []
    text = "\n".join(line for line in lines if line.strip())
    if not text.strip():
        warnings.append("Word 文件中没有文字（图片和公式对象不会被提取）")
    return Extracted(_limit(text, rules, warnings), None, tuple(warnings))


def decode_text(data: bytes, rules: UploadRules) -> Extracted:
    from charset_normalizer import from_bytes

    for bom, encoding in (
        (b"\xef\xbb\xbf", "utf-8-sig"),
        (b"\xff\xfe", "utf-16"),
        (b"\xfe\xff", "utf-16"),
    ):
        if data.startswith(bom):
            text = data.decode(encoding, errors="replace")
            break
    else:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            best = from_bytes(data).best()
            if best is None:
                raise UploadError("无法识别文本文件的编码") from None
            text = str(best)
    warnings: list[str] = []
    return Extracted(_limit(text.replace("\r\n", "\n"), rules, warnings), None, tuple(warnings))


def extract(file_type: FileType, data: bytes, rules: UploadRules) -> Extracted | None:
    """文档类直接提取文字；图片和音频需要模型处理，返回 None。"""
    if file_type.kind == "pdf":
        return extract_pdf(data, rules)
    if file_type.kind == "docx":
        return extract_docx(data, rules)
    if file_type.kind == "text":
        return decode_text(data, rules)
    return None
