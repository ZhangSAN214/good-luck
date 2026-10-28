"""Markdown 块 → Word（.docx），用已有的 python-docx 依赖。"""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from roundtable.core.export.markdown import Block, parse, runs

FONT = "Microsoft YaHei"
MONO = "Consolas"


def _east_asian(style_or_run, name: str) -> None:
    """中文字体要单独设 eastAsia，否则 Word 用默认的宋体。"""
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)


def _fill(paragraph, text: str, size: float | None = None) -> None:
    for r in runs(text):
        run = paragraph.add_run(r.text)
        run.bold = r.bold or None
        run.italic = r.italic or None
        if r.code:
            run.font.name = MONO
        if size:
            run.font.size = Pt(size)


def render(md: str) -> bytes:
    doc = DocxDocument()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    _east_asian(normal, FONT)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = doc.styles[name]
        style.font.name = FONT
        _east_asian(style, FONT)
    blocks: list[Block] = parse(md)
    for b in blocks:
        if b.kind == "heading":
            level = 0 if b.level == 1 else min(b.level - 1, 4)
            _fill(doc.add_heading("", level=level), b.text)
        elif b.kind == "bullet":
            style = "List Number" if b.ordered else "List Bullet"
            if b.level:
                style += f" {min(b.level + 1, 3)}"
            p = doc.add_paragraph(style=style)
            _fill(p, b.text)
        elif b.kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(b.text)
            run.font.name = MONO
            run.font.size = Pt(9.5)
        elif b.kind == "quote":
            p = doc.add_paragraph(style="Quote")
            _fill(p, b.text)
        elif b.kind == "table":
            table = doc.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.style = "Table Grid"
            for i, row in enumerate(b.rows):
                for j, cell in enumerate(row):
                    target = table.cell(i, j).paragraphs[0]
                    _fill(target, cell)
                    if i == 0:
                        for run in target.runs:
                            run.bold = True
            doc.add_paragraph()
        elif b.kind == "hr":
            p = doc.add_paragraph("—" * 20)
            p.runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        else:
            _fill(doc.add_paragraph(), b.text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
